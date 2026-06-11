import pandas as pd
from openpyxl import load_workbook
import sys

workbook = load_workbook(
    "C:/Users/arkya.mitra/OneDrive - OneWorkplace/Projects/FIFA_World_Cup_Predictor/data/worldcup-soccer-2026.xlsx",
    data_only=True
)


sheet = workbook["Round of 16,8,4, semi, final"]

for row in sheet.iter_rows(
    min_row=1,
    max_row=80,
    values_only=True
):
    print(row)

sys.exit()

groups_data = []

group_blocks = [

    {
        "group_row": 3,
        "team_rows": range(4,8)
    },

    {
        "group_row": 45,
        "team_rows": range(46,50)
    }

]

group_columns = [2,5,8,11,14,17]

for block in group_blocks:

    for column in group_columns:

        group_name = sheet.cell(
            row=block["group_row"],
            column=column
        ).value

        for row in block["team_rows"]:

            team = sheet.cell(
                row=row,
                column=column
            ).value

            groups_data.append({

                "Group": group_name,
                "Team": team

            })


group_df = pd.DataFrame(
    groups_data
)

group_df["Group"] = group_df["Group"].replace(
    {
        "K": "Group I"
    }
)

fixtures_data = []

group_columns = [

    (2,3,"Group A"),
    (5,6,"Group B"),
    (8,9,"Group C"),
    (11,12,"Group D"),
    (14,15,"Group E"),
    (17,18,"Group F"),

    (2,3,"Group G"),
    (5,6,"Group H"),
    (8,9,"Group I"),
    (11,12,"Group J"),
    (14,15,"Group K"),
    (17,18,"Group L")

]

fixture_rows = [

    [11,16,21,26,31,36],   # Groups A-F

    [53,58,63,68,73,78]    # Groups G-L

]

for i, (home_col, away_col, group_name) in enumerate(group_columns):

    if i < 6:

        rows = fixture_rows[0]

    else:

        rows = fixture_rows[1]

    for row in rows:

        home_team = sheet.cell(
            row= row,
            column= home_col
        ).value

        away_team = sheet.cell(
            row= row,
            column= away_col
        ).value

        fixtures_data.append({
            "Group": group_name,
            "Home Team": home_team,
            "Away_Team": away_team
        })



fixtures_df = pd.DataFrame(
    fixtures_data
)

print(fixtures_df.head(10))


group_df.to_csv(

    "C:/Users/arkya.mitra/OneDrive - OneWorkplace/Projects/FIFA_World_Cup_Predictor/data/world_cup_groups.csv",

    index=False
)

print(
    "Total teams:",
    len(group_df)
)

print(
    "Total fixtures:",
    len(fixtures_df)
)

fixtures_df.to_csv(

    "C:/Users/arkya.mitra/OneDrive - OneWorkplace/Projects/FIFA_World_Cup_Predictor/data/world_cup_fixtures.csv",

    index=False
)